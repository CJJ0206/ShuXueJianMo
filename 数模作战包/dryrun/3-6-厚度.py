# throwaway: 今晚压测用，不作任何竞赛结论
"""A 路：去包络（减 12 阶多项式基线）→ 等间隔重采样 → 加窗 FFT → 条纹频率 f → 厚度 d。

物理关系（今晚第一次写错过，务必记住方向）：
    干涉项 ~ cos(2π · 2nd · ν̃)，ν̃ 为波数 (cm⁻¹)
    ⇒ 条纹频率 f [cycle per cm⁻¹] = 2·n·d [cm]
    ⇒ d = f / (2n)          ← 不是 1/(2nf)！写反后两个错误会互相掩盖。
先验区间 d∈[1,15] μm ⇒ f∈[5.2e-4, 7.8e-3] ⇒ bin∈[1.9, 28]。
n(SiC)=2.6 为占位，明天换成题面值并登记 C##。"""
import json
import time
from pathlib import Path

import numpy as np
import openpyxl

SRC = "C:/Users/Administrator/Desktop/CUMCM2025Problems/B题/附件/附件1.xlsx"
OUT = Path(__file__).with_name("3-6-result-A.json")
N_SIC = 2.6
D_LO_UM, D_HI_UM = 1.0, 15.0


def load(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    arr = np.array([r[:2] for r in ws.iter_rows(min_row=2, values_only=True)], dtype=float)
    wb.close()
    wn, ref = arr[:, 0], arr[:, 1]
    m = np.isfinite(wn) & np.isfinite(ref)
    o = np.argsort(wn[m])
    return wn[m][o], ref[m][o]


def main():
    wn, ref = load(SRC)
    wn, ref = wn[1:], ref[1:]                       # 剔除首行 0（假设，B 路不剔）
    npt = len(wn)
    uni = np.linspace(wn[0], wn[-1], npt)
    y = np.interp(uni, wn, ref)
    res = 1.0 / (uni[-1] - uni[0])                  # 频率分辨单元 cycle/cm⁻¹

    det = y - np.polyval(np.polyfit(uni, y, 12), uni)     # 去包络（B 路不做）
    spec = np.abs(np.fft.rfft((det - det.mean()) * np.hanning(npt)))

    lo, hi = 2 * N_SIC * D_LO_UM * 1e-4, 2 * N_SIC * D_HI_UM * 1e-4
    blo, bhi = int(np.ceil(lo / res)), int(np.floor(hi / res))
    band = spec[blo:bhi + 1]
    kbin = int(np.argmax(band)) + blo
    freq = kbin * res
    d_um = freq / (2 * N_SIC) * 1e4                  # cm → μm

    print(f"[A 路] 分辨单元={res:.3e} 先验带 bin∈[{blo},{bhi}] 选出 bin={kbin} "
          f"f={freq:.6f} cm⁻¹  厚度={d_um:.2f} μm")
    assert blo <= kbin <= bhi, "选峰落在先验带外，说明去包络失败"
    OUT.write_text(json.dumps({"path": "detrend+FFT", "src": SRC, "kbin": kbin,
                               "res_per_cm": res, "freq_per_cm": freq, "thickness_um": d_um,
                               "n_sic": N_SIC, "band_bins": [blo, bhi],
                               "assumptions": ["去 12 阶多项式包络", "重采样等间隔", "剔除首行0",
                                               "单基频正弦"],
                               "ts": time.strftime("%F %T")}, ensure_ascii=False, indent=2),
                   encoding="utf-8")


if __name__ == "__main__":
    main()
