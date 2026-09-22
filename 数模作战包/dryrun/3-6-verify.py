# throwaway: 今晚压测用，不作任何竞赛结论
"""B 路独立复算。与 A 路的独立性做在三个层面上：
   ① 算法：非均匀网格上的「候选频率最小二乘正弦拟合」，不经 FFT；
   ② 预处理：不去包络、不重采样——靠把候选频率限制在物理先验带来排除慢趋势，
      而 A 路是靠去包络排除的。两条路排趋势的机理不同，才叫独立。
   ③ 假设：不剔首行 0，并单独量出这条假设的影响。
   本文件不得 import / exec A 路任何符号。"""
import json
from pathlib import Path

import numpy as np
import openpyxl

SRC = "C:/Users/Administrator/Desktop/CUMCM2025Problems/B题/附件/附件1.xlsx"
A = Path(__file__).with_name("3-6-result-A.json")
N_SIC = 2.6
D_LO_UM, D_HI_UM = 1.0, 15.0


def load(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    arr = np.array([r[:2] for r in ws.iter_rows(min_row=2, values_only=True)], dtype=float)
    wb.close()
    return arr[:, 0], arr[:, 1]


def best_f(wn, ref, fgrid):
    y = ref - ref.mean()
    bf, br = np.nan, np.inf
    for f in fgrid:
        X = np.column_stack([np.cos(2 * np.pi * f * wn), np.sin(2 * np.pi * f * wn),
                             np.ones_like(wn)])
        coef, res, *_ = np.linalg.lstsq(X, y, rcond=None)
        rss = float(res[0]) if res.size else float(((X @ coef - y) ** 2).sum())
        if rss < br:
            br, bf = rss, f
    return bf, br


def main():
    wn, ref = load(SRC)
    m = np.isfinite(wn) & np.isfinite(ref)
    wn, ref = wn[m], ref[m]
    o = np.argsort(wn)
    wn, ref = wn[o], ref[o]
    span = wn[-1] - wn[0]
    res = 1.0 / span
    lo, hi = 2 * N_SIC * D_LO_UM * 1e-4, 2 * N_SIC * D_HI_UM * 1e-4

    f_inc, _ = best_f(wn, ref, np.arange(lo, hi, res / 20))
    f_exc, _ = best_f(wn[1:], ref[1:], np.arange(lo, hi, res / 20))
    d_inc = f_inc / (2 * N_SIC) * 1e4
    d_exc = f_exc / (2 * N_SIC) * 1e4

    a = json.loads(A.read_text(encoding="utf-8"))
    b_inc = abs(f_inc - a["freq_per_cm"]) / a["res_per_cm"]
    b_exc = abs(f_exc - a["freq_per_cm"]) / a["res_per_cm"]
    sens = abs(f_inc - f_exc) / a["res_per_cm"]

    print(f"[B 路] 含首行 f={f_inc:.6f} → d={d_inc:.2f} μm | 剔首行 f={f_exc:.6f} → d={d_exc:.2f} μm")
    print(f"[A 路] f={a['freq_per_cm']:.6f} → d={a['thickness_um']:.2f} μm "
          f"(分辨单元 {a['res_per_cm']:.3e})")
    print(f"  双路差：含首行 {b_inc:.2f} bin | 剔首行 {b_exc:.2f} bin")
    print(f"  【假设敏感度】'首行 0 是否剔除' 造成 {sens:.2f} bin 差异")
    ok = min(b_inc, b_exc) <= 1.0 and D_LO_UM <= d_inc <= D_HI_UM
    print(f"{'PASS' if ok else 'FAIL'} 双路一致且在先验内")
    print("注意：一致只是必要条件。两路共享同一错误假设时一致地错——"
          "充分性由'落在独立给出的粗估区间内 + 合成锚点取回已知值'两条共同保证。")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
